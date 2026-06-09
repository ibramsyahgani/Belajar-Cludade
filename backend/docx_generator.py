from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from .models import CVData
from .config import load_config


def hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def add_horizontal_rule(paragraph):
    """Add a thin bottom border to a paragraph (section divider)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_paragraph_spacing(para, before=0, after=2):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def generate_docx(cv_data: CVData) -> bytes:
    config = load_config()
    style_cfg = config.get("docx_style", {})
    palettes = config.get("color_palettes", {})
    headings_cfg = config.get("section_headings", {})
    section_order = config.get("section_order", [
        "ringkasan_profesional", "pendidikan", "pengalaman",
        "sertifikasi_pencapaian", "leadership_activities", "kemampuan"
    ])

    font_name = style_cfg.get("font_name", "Calibri")
    font_size = style_cfg.get("font_size", 10)
    name_size = style_cfg.get("name_size", 16)
    heading_size = style_cfg.get("heading_size", 11)
    margin_cm = style_cfg.get("margin_cm", 2.5)

    palette = palettes.get(cv_data.bidang, palettes.get("default", {"primary": "#1F3A5F", "accent": "#4A6FA5"}))
    primary_rgb = hex_to_rgb(palette["primary"])

    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(margin_cm)
        section.bottom_margin = Cm(margin_cm)
        section.left_margin = Cm(margin_cm)
        section.right_margin = Cm(margin_cm)

    # --- HEADER ---
    h = cv_data.header
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name_para, 0, 2)
    run = name_para.add_run(h.nama or "")
    run.bold = True
    run.font.size = Pt(name_size)
    run.font.name = font_name

    # Contact line
    contact_parts = [p for p in [h.lokasi, h.email, h.telepon, h.linkedin_atau_portfolio] if p]
    if contact_parts:
        contact_para = doc.add_paragraph(" | ".join(contact_parts))
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(contact_para, 0, 6)
        for run in contact_para.runs:
            run.font.size = Pt(font_size)
            run.font.name = font_name
            run.font.color.rgb = RGBColor(100, 100, 100)

    def add_section_heading(title: str):
        para = doc.add_paragraph()
        set_paragraph_spacing(para, 8, 2)
        run = para.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(heading_size)
        run.font.name = font_name
        run.font.color.rgb = RGBColor(*primary_rgb)
        add_horizontal_rule(para)
        return para

    def add_entry_header(org: str, loc: str, period: str, role: str = ""):
        # Org + period on same line (org bold left, period right)
        para = doc.add_paragraph()
        set_paragraph_spacing(para, 4, 0)
        run_org = para.add_run(org)
        run_org.bold = True
        run_org.font.size = Pt(font_size)
        run_org.font.name = font_name
        if loc:
            run_loc = para.add_run(f", {loc}")
            run_loc.font.size = Pt(font_size)
            run_loc.font.name = font_name
        # Right-align period using tab stop
        if period:
            tab_run = para.add_run(f"\t{period}")
            tab_run.font.size = Pt(font_size)
            tab_run.font.name = font_name
            para.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_ALIGN_PARAGRAPH.RIGHT)

        if role:
            role_para = doc.add_paragraph()
            set_paragraph_spacing(role_para, 0, 0)
            role_run = role_para.add_run(role)
            role_run.italic = True
            role_run.font.size = Pt(font_size)
            role_run.font.name = font_name

    def add_bullet(text: str):
        para = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(para, 0, 1)
        para.paragraph_format.left_indent = Cm(0.5)
        run = para.add_run(text)
        run.font.size = Pt(font_size)
        run.font.name = font_name

    def add_normal(text: str, italic=False, size=None):
        para = doc.add_paragraph()
        set_paragraph_spacing(para, 0, 2)
        run = para.add_run(text)
        run.font.size = Pt(size or font_size)
        run.font.name = font_name
        run.italic = italic
        return para

    # Render sections in configured order
    for section_key in section_order:
        if section_key == "ringkasan_profesional" and cv_data.ringkasan_profesional:
            add_section_heading(headings_cfg.get("ringkasan_profesional", "RINGKASAN PROFESIONAL"))
            add_normal(cv_data.ringkasan_profesional)

        elif section_key == "pendidikan" and cv_data.pendidikan:
            add_section_heading(headings_cfg.get("pendidikan", "PENDIDIKAN & PELATIHAN"))
            for edu in cv_data.pendidikan:
                add_entry_header(edu.institusi, edu.lokasi, edu.periode, edu.jurusan_gelar)
                if edu.ipk:
                    add_normal(f"IPK: {edu.ipk}")
                if edu.skripsi:
                    add_normal(f"Skripsi: {edu.skripsi}", italic=True)
                for pel in edu.pelatihan:
                    if pel:
                        add_bullet(pel)

        elif section_key == "pengalaman" and cv_data.pengalaman:
            add_section_heading(headings_cfg.get("pengalaman", "PENGALAMAN"))
            for exp in cv_data.pengalaman:
                add_entry_header(exp.organisasi, exp.lokasi, exp.periode, exp.posisi)
                for poin in exp.poin:
                    if poin:
                        add_bullet(poin)

        elif section_key == "sertifikasi_pencapaian" and cv_data.sertifikasi_pencapaian:
            add_section_heading(headings_cfg.get("sertifikasi_pencapaian", "SERTIFIKASI & PENCAPAIAN"))
            for item in cv_data.sertifikasi_pencapaian:
                if item:
                    add_bullet(item)

        elif section_key == "leadership_activities" and cv_data.leadership_activities:
            add_section_heading(headings_cfg.get("leadership_activities", "LEADERSHIP & AKTIVITAS"))
            for act in cv_data.leadership_activities:
                add_entry_header(act.organisasi, act.lokasi, act.periode, act.peran)
                for poin in act.poin:
                    if poin:
                        add_bullet(poin)

        elif section_key == "kemampuan" and cv_data.kemampuan:
            add_section_heading(headings_cfg.get("kemampuan", "KEMAMPUAN"))
            k = cv_data.kemampuan
            if k.hard_skills:
                para = doc.add_paragraph()
                set_paragraph_spacing(para, 0, 1)
                r = para.add_run("Hard Skills: ")
                r.bold = True
                r.font.size = Pt(font_size)
                r.font.name = font_name
                r2 = para.add_run(", ".join(k.hard_skills))
                r2.font.size = Pt(font_size)
                r2.font.name = font_name
            if k.tools_software:
                para = doc.add_paragraph()
                set_paragraph_spacing(para, 0, 1)
                r = para.add_run("Tools & Software: ")
                r.bold = True
                r.font.size = Pt(font_size)
                r.font.name = font_name
                r2 = para.add_run(", ".join(k.tools_software))
                r2.font.size = Pt(font_size)
                r2.font.name = font_name
            if k.bahasa:
                para = doc.add_paragraph()
                set_paragraph_spacing(para, 0, 1)
                r = para.add_run("Bahasa: ")
                r.bold = True
                r.font.size = Pt(font_size)
                r.font.name = font_name
                r2 = para.add_run(", ".join(k.bahasa))
                r2.font.size = Pt(font_size)
                r2.font.name = font_name

    # Notes section (always at end if present)
    if cv_data.notes:
        add_section_heading("CATATAN UNTUK ANDA")
        add_normal(cv_data.notes)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
